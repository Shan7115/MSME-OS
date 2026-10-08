import os
import io
import csv
from typing import List, Dict, Any, Tuple

# Supported formats
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".txt", ".csv"}
MAX_FILE_SIZE = 25 * 1024 * 1024  # 25 MB

class DocumentExtractionError(Exception):
    pass

def extract_pdf(file_bytes: bytes) -> Tuple[int, List[Dict[str, Any]], str]:
    import pymupdf
    pages = []
    full_text_parts = []
    
    with pymupdf.open(stream=file_bytes, filetype="pdf") as doc:
        page_count = len(doc)
        for page_idx in range(page_count):
            page = doc[page_idx]
            page_text = page.get_text("text").strip()
            
            # Extract basic table/block information if available
            blocks = page.get_text("blocks")
            pages.append({
                "page_number": page_idx + 1,
                "text": page_text,
                "block_count": len(blocks)
            })
            if page_text:
                full_text_parts.append(f"--- PAGE {page_idx + 1} ---\n{page_text}")
                
    return page_count, pages, "\n\n".join(full_text_parts)

def extract_docx(file_bytes: bytes) -> Tuple[int, List[Dict[str, Any]], str]:
    import docx
    doc = docx.Document(io.BytesIO(file_bytes))
    
    text_blocks = []
    for p in doc.paragraphs:
        if p.text.strip():
            text_blocks.append(p.text.strip())
            
    for table in doc.tables:
        table_rows = []
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells]
            if any(row_cells):
                table_rows.append(" | ".join(row_cells))
        if table_rows:
            text_blocks.append("\n[Table Data]:\n" + "\n".join(table_rows))
            
    full_text = "\n\n".join(text_blocks)
    # Estimate page count for docx (~400 words per page)
    word_count = len(full_text.split())
    page_count = max(1, (word_count + 399) // 400)
    
    pages = [{
        "page_number": 1,
        "text": full_text,
        "block_count": len(text_blocks)
    }]
    return page_count, pages, full_text

def extract_xlsx(file_bytes: bytes) -> Tuple[int, List[Dict[str, Any]], str]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    sheet_texts = []
    pages = []
    
    for idx, sheetname in enumerate(wb.sheetnames):
        ws = wb[sheetname]
        rows_data = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip() != ""]
            if cells:
                rows_data.append(" | ".join(cells))
                
        sheet_text = f"Sheet: {sheetname}\n" + "\n".join(rows_data)
        sheet_texts.append(sheet_text)
        pages.append({
            "page_number": idx + 1,
            "sheet_name": sheetname,
            "text": sheet_text
        })
        
    full_text = "\n\n".join(sheet_texts)
    return len(wb.sheetnames), pages, full_text

def extract_txt(file_bytes: bytes) -> Tuple[int, List[Dict[str, Any]], str]:
    try:
        text = file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        text = file_bytes.decode("latin-1", errors="replace")
        
    pages = [{"page_number": 1, "text": text}]
    return 1, pages, text

def extract_csv(file_bytes: bytes) -> Tuple[int, List[Dict[str, Any]], str]:
    try:
        decoded = file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        decoded = file_bytes.decode("latin-1", errors="replace")
        
    reader = csv.reader(io.StringIO(decoded))
    lines = []
    for row in reader:
        if row:
            lines.append(" | ".join(row))
    full_text = "\n".join(lines)
    pages = [{"page_number": 1, "text": full_text}]
    return 1, pages, full_text

def extract_document(filename: str, file_bytes: bytes) -> Dict[str, Any]:
    if len(file_bytes) > MAX_FILE_SIZE:
        raise DocumentExtractionError(f"File size exceeds maximum allowable limit of {MAX_FILE_SIZE // (1024*1024)} MB")
        
    ext = os.path.splitext(filename)[1].lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise DocumentExtractionError(f"Unsupported file type '{ext}'. Supported formats: {', '.join(SUPPORTED_EXTENSIONS)}")
        
    if len(file_bytes) == 0:
        raise DocumentExtractionError("Uploaded file is empty (0 bytes).")
        
    if ext == ".pdf":
        page_count, pages, full_text = extract_pdf(file_bytes)
    elif ext == ".docx":
        page_count, pages, full_text = extract_docx(file_bytes)
    elif ext == ".xlsx":
        page_count, pages, full_text = extract_xlsx(file_bytes)
    elif ext == ".txt":
        page_count, pages, full_text = extract_txt(file_bytes)
    elif ext == ".csv":
        page_count, pages, full_text = extract_csv(file_bytes)
    else:
        raise DocumentExtractionError(f"Handler for extension '{ext}' not implemented.")
        
    if not full_text.strip():
        raise DocumentExtractionError("No readable text could be extracted from the document.")
        
    return {
        "filename": filename,
        "file_type": ext,
        "file_size": len(file_bytes),
        "page_count": page_count,
        "pages": pages,
        "extracted_text": full_text
    }
